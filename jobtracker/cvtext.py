"""Turn an uploaded CV (PDF, Word, text) into plain text, and split it into evidence-sized pieces."""

from __future__ import annotations

import io
import re

SECTION_WORDS = {
    "experience", "work experience", "professional experience", "employment", "research experience", "education",
    "skills", "technical skills", "publications", "projects", "certifications", "languages", "awards", "summary",
    "profile", "volunteering", "leadership", "activities", "teaching", "presentations", "interests", "courses",
}
BULLET = re.compile(r"^\s*(?:[•▪◦●○■□➢►▶–—\-*·]|\d+[.)])\s+")


def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        text = "\n".join((p.extract_text() or "") for p in reader.pages)
    elif name.endswith(".docx"):
        import docx

        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        text = "\n".join(parts)
    elif name.endswith((".txt", ".md")):
        text = data.decode("utf-8", errors="replace")
    else:
        raise ValueError("Upload a PDF, Word (.docx) or text file.")
    text = text.replace(" ", " ").replace("\r", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    if len(text.strip()) < 50:
        raise ValueError("Couldn't read any text from this file. If it's a scanned PDF, paste the text instead.")
    return text.strip()


def chunks(cv: str) -> list[dict]:
    """Split a CV into bullet/sentence pieces with the section they belong to.

    PDF text often breaks one bullet over several lines, so lines are joined until the next bullet,
    heading or blank line.
    """
    pieces: list[dict] = []
    section = "Other"
    buf: list[str] = []

    def flush():
        if buf:
            text = " ".join(buf).strip()
            for sent in re.split(r"(?<=[.;])\s+(?=[A-Z])", text):
                s = sent.strip(" ;")
                if len(s) >= 18:
                    pieces.append({"text": s, "section": section})
            buf.clear()

    for raw in cv.splitlines():
        line = raw.strip()
        if not line:
            flush()
            continue
        bare = line.rstrip(":").strip().lower()
        if bare in SECTION_WORDS or (line.isupper() and 3 <= len(line) <= 40 and len(line.split()) <= 4):
            flush()
            section = line.rstrip(":").strip().title()
            continue
        if BULLET.match(line):
            flush()
            buf.append(BULLET.sub("", line))
        else:
            buf.append(line)
    flush()
    return pieces
